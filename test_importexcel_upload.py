from __future__ import annotations

import asyncio
import os
import shutil
import sqlite3
import tempfile
import unittest
from contextlib import closing
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook

import telegram_license_bot as bot
from scripts.import_inventory import REQUIRED_COLUMNS


PROJECT_ROOT = Path(__file__).resolve().parent


class FakeTelegramFile:
    def __init__(self, source_path: Path) -> None:
        self.source_path = source_path

    async def download_to_drive(self, custom_path: str) -> None:
        shutil.copy2(self.source_path, custom_path)


class FakeDocument:
    def __init__(self, source_path: Path, file_name: str, file_size: int | None = None) -> None:
        self.source_path = source_path
        self.file_name = file_name
        self.file_size = file_size if file_size is not None else (source_path.stat().st_size if source_path.exists() else 0)

    async def get_file(self) -> FakeTelegramFile:
        return FakeTelegramFile(self.source_path)


class FakeMessage:
    def __init__(self, document: FakeDocument) -> None:
        self.document = document
        self.replies: list[str] = []

    async def reply_text(self, text: str, reply_markup=None) -> None:
        self.replies.append(text)


class ImportExcelUploadTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.db_path = self.root / "store.db"
        shutil.copy2(PROJECT_ROOT / "database" / "store.db", self.db_path)
        self.workbook_path = self.root / "capcut_7d_upload.xlsx"
        self._write_workbook(self.workbook_path)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _write_workbook(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CAPCUT"
        sheet.append(list(REQUIRED_COLUMNS))
        sheet.append([
            "CAPCUT",
            "AI",
            "CAPCUT",
            "personal",
            "7D",
            8000,
            7,
            "upload-capcut7d-one@example.com|secret-pass",
            "",
            1,
        ])
        workbook.save(path)

    def _write_capcut_30d_workbook(self, path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CAPCUT"
        sheet.append(list(REQUIRED_COLUMNS))
        sheet.append([
            "CAPCUT",
            "AI",
            "CAPCUT PRO",
            "personal",
            "30D",
            45000,
            30,
            "upload-capcut30d-one@example.com|secret-pass",
            "",
            1,
        ])
        workbook.save(path)

    def _context(self, *, running: bool = False):
        return SimpleNamespace(
            application=SimpleNamespace(
                bot_data={
                    "admin_ids": {99},
                    "store_db_path": self.db_path,
                    "inventory_import_running": running,
                }
            )
        )

    def _update(self, document: FakeDocument, *, user_id: int = 99):
        message = FakeMessage(document)
        update = SimpleNamespace(
            effective_user=SimpleNamespace(id=user_id),
            effective_message=message,
        )
        return update, message

    def test_admin_document_upload_imports_capcut_7d_and_masks_secret(self) -> None:
        document = FakeDocument(self.workbook_path, "capcut_7d_upload.xlsx")
        update, message = self._update(document)
        context = self._context()

        asyncio.run(bot.on_import_document(update, context))

        self.assertEqual(len(message.replies), 1)
        reply = message.replies[0]
        self.assertIn("Product code: CAPCUT_7D", reply)
        self.assertIn("Stock", reply)
        self.assertIn("Account", reply)
        self.assertIn(": 1", reply)
        self.assertIn("Tồn kho CAPCUT_7D sau import: 1", reply)
        self.assertNotIn("secret-pass", reply)
        self.assertNotIn("upload-capcut7d-one@example.com", reply)

        with closing(sqlite3.connect(self.db_path)) as connection:
            row = connection.execute(
                """
                SELECT p.code, i.status, COUNT(*)
                FROM inventory_items AS i
                JOIN products AS p ON p.id = i.product_id
                WHERE p.code = 'CAPCUT_7D'
                GROUP BY p.code, i.status
                """
            ).fetchone()
        self.assertEqual(row, ("CAPCUT_7D", "available", 1))

    def test_admin_document_upload_maps_capcut_30d_terms_to_capcut_30d_sku(self) -> None:
        workbook_path = self.root / "capcut_30d_upload.xlsx"
        self._write_capcut_30d_workbook(workbook_path)
        document = FakeDocument(workbook_path, "capcut_30d_upload.xlsx")
        update, message = self._update(document)

        asyncio.run(bot.on_import_document(update, self._context()))

        self.assertEqual(len(message.replies), 1)
        reply = message.replies[0]
        self.assertIn("Product code: CAPCUT_30D", reply)
        self.assertIn("CAPCUT_30D sau import: 1", reply)
        self.assertNotIn("CAPCUT requires price_vnd=400000", reply)
        self.assertNotIn("secret-pass", reply)

        with closing(sqlite3.connect(self.db_path)) as connection:
            row = connection.execute(
                """
                SELECT p.code, i.status, COUNT(*)
                FROM inventory_items AS i
                JOIN products AS p ON p.id = i.product_id
                WHERE p.code = 'CAPCUT_30D'
                GROUP BY p.code, i.status
                """
            ).fetchone()
        self.assertEqual(row, ("CAPCUT_30D", "available", 1))

    def test_non_admin_document_upload_is_rejected(self) -> None:
        document = FakeDocument(self.workbook_path, "capcut_7d_upload.xlsx")
        update, message = self._update(document, user_id=100)
        context = self._context()

        asyncio.run(bot.on_import_document(update, context))

        self.assertEqual(message.replies, ["Bạn không có quyền thực hiện thao tác này."])

    def test_wrong_file_type_is_rejected(self) -> None:
        text_path = self.root / "not_inventory.txt"
        text_path.write_text("not an inventory file", encoding="utf-8")
        update, message = self._update(FakeDocument(text_path, "not_inventory.txt"))

        asyncio.run(bot.on_import_document(update, self._context()))

        self.assertEqual(message.replies, ["File import phải có định dạng .xlsx hoặc .csv."])

    def test_empty_file_is_rejected(self) -> None:
        empty_path = self.root / "empty.xlsx"
        empty_path.write_bytes(b"")
        update, message = self._update(FakeDocument(empty_path, "empty.xlsx"))

        asyncio.run(bot.on_import_document(update, self._context()))

        self.assertEqual(message.replies, ["File import đang rỗng hoặc không có dữ liệu."])

    def test_oversized_file_is_rejected(self) -> None:
        update, message = self._update(
            FakeDocument(self.workbook_path, "capcut_7d_upload.xlsx", file_size=10 * 1024 * 1024 + 1)
        )

        asyncio.run(bot.on_import_document(update, self._context()))

        self.assertEqual(message.replies, ["File import vượt quá giới hạn 10 MB."])

    def test_corrupt_file_returns_safe_error(self) -> None:
        corrupt_path = self.root / "corrupt.xlsx"
        corrupt_path.write_text("not a real xlsx with secret-pass", encoding="utf-8")
        update, message = self._update(FakeDocument(corrupt_path, "corrupt.xlsx"))

        asyncio.run(bot.on_import_document(update, self._context()))

        self.assertEqual(len(message.replies), 1)
        self.assertIn("Import SQLite thất bại:", message.replies[0])
        self.assertNotIn("secret-pass", message.replies[0])

    def test_duplicate_upload_is_reported_without_adding_stock(self) -> None:
        first_update, first_message = self._update(FakeDocument(self.workbook_path, "capcut_7d_upload.xlsx"))
        asyncio.run(bot.on_import_document(first_update, self._context()))
        self.assertIn("Tồn kho CAPCUT_7D sau import: 1", first_message.replies[0])

        second_update, second_message = self._update(FakeDocument(self.workbook_path, "capcut_7d_upload.xlsx"))
        asyncio.run(bot.on_import_document(second_update, self._context()))

        reply = second_message.replies[0]
        self.assertIn("Account", reply)
        self.assertIn(": 0", reply)
        self.assertIn("1", reply)
        self.assertIn("Stock sau import: 1", reply)

    def test_concurrent_import_is_rejected(self) -> None:
        update, message = self._update(FakeDocument(self.workbook_path, "capcut_7d_upload.xlsx"))

        asyncio.run(bot.on_import_document(update, self._context(running=True)))

        self.assertEqual(message.replies, ["Đang có phiên import khác chạy. Vui lòng thử lại sau."])


if __name__ == "__main__":
    unittest.main(verbosity=2)
