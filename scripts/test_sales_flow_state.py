from __future__ import annotations

import os
import sqlite3
import tempfile
import unittest
from contextlib import closing
from datetime import datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
import sys

sys.path.insert(0, str(PROJECT_ROOT))

from repository.store_repository import ACCOUNT_PRODUCT_CODE_ALIASES, StoreRepository
from scripts.import_inventory import REQUIRED_COLUMNS, import_inventory
from scripts.sales_flow_state import (
    assert_reservation_transition,
    canonical_product_code,
    log_sales_state,
    snapshot_sales_state,
    validate_stock_invariant,
)
import telegram_license_bot as bot


class SalesFlowStateTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "store.db"
        source_db = PROJECT_ROOT / "database" / "store.db"
        with closing(sqlite3.connect(source_db)) as source, closing(sqlite3.connect(self.db_path)) as target:
            source.backup(target)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _clear_inventory_for_codes(self, *codes: str) -> None:
        normalized = {code.upper() for code in codes if code}
        expanded = set(normalized)
        for code in normalized:
            expanded.update(alias.upper() for alias in ACCOUNT_PRODUCT_CODE_ALIASES.get(code, ()))
        placeholders = ",".join("?" for _ in expanded)
        with closing(sqlite3.connect(self.db_path)) as connection, connection:
            connection.execute("PRAGMA foreign_keys = ON")
            item_ids = [
                row[0]
                for row in connection.execute(
                    f"""
                    SELECT i.id
                    FROM inventory_items AS i
                    JOIN products AS p ON p.id = i.product_id
                    WHERE UPPER(p.code) IN ({placeholders})
                    """,
                    tuple(sorted(expanded)),
                ).fetchall()
            ]
            if not item_ids:
                return
            item_placeholders = ",".join("?" for _ in item_ids)
            connection.execute(
                f"DELETE FROM inventory_movements WHERE inventory_item_id IN ({item_placeholders})",
                tuple(item_ids),
            )
            connection.execute(
                f"DELETE FROM order_inventory_items WHERE inventory_item_id IN ({item_placeholders})",
                tuple(item_ids),
            )
            connection.execute(
                f"DELETE FROM inventory_items WHERE id IN ({item_placeholders})",
                tuple(item_ids),
            )

    def _import_workbook(self, sheet_name: str, credentials: list[str], *, price_vnd: int = 70000, warranty_days: int = 7, mode: str = "replace") -> dict[str, object]:
        workbook_path = Path(self.temp_dir.name) / f"{sheet_name.lower()}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(REQUIRED_COLUMNS)
        for credential in credentials:
            sheet.append([sheet_name, "account", f"{sheet_name} PRO", "account", "30D", price_vnd, warranty_days, credential, "", 1])
        workbook.save(workbook_path)
        workbook.close()
        return import_inventory(workbook_path, self.db_path, mode=mode)

    def _reserve_gemini_quantity(self, quantity: int, *, order_id: str) -> dict[str, object]:
        previous_store_db_path = os.environ.get("STORE_DB_PATH")
        previous_make_order_id = bot._make_order_id
        try:
            os.environ["STORE_DB_PATH"] = str(self.db_path)
            bot._make_order_id = lambda _product_name: order_id
            fake_update = type(
                "FakeOrderUpdate",
                (),
                {"effective_user": type("FakeUser", (), {"id": 42, "full_name": "Test User", "username": ""})()},
            )()
            return bot._create_sales_order(fake_update, "GEMINI AI", "GEMINI", quantity)
        finally:
            bot._make_order_id = previous_make_order_id
            if previous_store_db_path is None:
                os.environ.pop("STORE_DB_PATH", None)
            else:
                os.environ["STORE_DB_PATH"] = previous_store_db_path

    def test_import_gemini_8_and_duplicate_credentials_keep_slots(self) -> None:
        report = self._import_workbook("GEMINI", ["gem1@example.com|pass", "gem2@example.com|pass", "gem2@example.com|pass", "gem4@example.com|pass", "gem5@example.com|pass", "gem6@example.com|pass", "gem7@example.com|pass", "gem8@example.com|pass"])
        self.assertEqual(report["stock"], {"GEMINI": 8})
        state = snapshot_sales_state(self.db_path, "GEMINI")
        self.assertEqual(state.available_count, 8)
        self.assertEqual(state.reserved_count, 0)
        self.assertEqual(state.delivered_count, 0)
        self.assertEqual(state.disabled_count, 0)
        validate_stock_invariant(state)

    def test_quantity_boundary_reserve_success_and_failures(self) -> None:
        self._clear_inventory_for_codes("GEMINI")
        self._import_workbook(
            "GEMINI",
            [f"gem{i}@example.com|pass" for i in range(1, 9)],
            price_vnd=70000,
            warranty_days=7,
        )
        package = bot._get_package_info("GEMINI AI", "GEMINI")
        self.assertIsNotNone(package)
        self.assertEqual(canonical_product_code(str(package["product_code"])), "GEMINI")
        self.assertEqual(int(package["available_count"]), 8)

        before = snapshot_sales_state(self.db_path, "GEMINI", callback_data="qty:GEMINI AI:GEMINI:1", package_id="GEMINI", quantity=1)
        self.assertTrue(before.can_reserve)
        self.assertEqual(before.expected_available_after, 7)
        self.assertEqual(before.expected_reserved_after, 1)

        order = self._reserve_gemini_quantity(1, order_id="ORD-GEMINI-1")
        self.assertEqual(order["product_code"], "GEMINI")
        after = snapshot_sales_state(self.db_path, "GEMINI", order_id="ORD-GEMINI-1", callback_data="qty:GEMINI AI:GEMINI:1", package_id="GEMINI", quantity=1)
        self.assertEqual(after.available_count, 7)
        self.assertEqual(after.reserved_count, 1)
        self.assertEqual(after.reserved_item_ids, tuple(after.reserved_item_ids))
        assert_reservation_transition(before, after, 1)

        self._clear_inventory_for_codes("GEMINI")
        self._import_workbook("GEMINI", [f"gem{i}@example.com|pass" for i in range(1, 9)], price_vnd=70000, warranty_days=7)
        order = self._reserve_gemini_quantity(8, order_id="ORD-GEMINI-8")
        self.assertEqual(order["product_code"], "GEMINI")
        state = snapshot_sales_state(self.db_path, "GEMINI", order_id="ORD-GEMINI-8", callback_data="qty:GEMINI AI:GEMINI:8", package_id="GEMINI", quantity=8)
        self.assertEqual(state.available_count, 0)
        self.assertEqual(state.reserved_count, 8)
        validate_stock_invariant(state)

        self._clear_inventory_for_codes("GEMINI")
        self._import_workbook("GEMINI", [f"gem{i}@example.com|pass" for i in range(1, 9)], price_vnd=70000, warranty_days=7)
        with self.assertRaisesRegex(bot.InventoryReservationError, "hết hàng"):
            self._reserve_gemini_quantity(9, order_id="ORD-GEMINI-9")
        state = snapshot_sales_state(self.db_path, "GEMINI", callback_data="qty:GEMINI AI:GEMINI:9", package_id="GEMINI", quantity=9)
        self.assertEqual(state.available_count, 8)
        self.assertEqual(state.reserved_count, 0)
        self.assertFalse(state.can_reserve)
        self.assertIn("available_before=8", state.reason)

    def test_timeout_paid_and_delivery_invariants(self) -> None:
        self._clear_inventory_for_codes("GEMINI")
        self._import_workbook("GEMINI", [f"gem{i}@example.com|pass" for i in range(1, 9)], price_vnd=70000, warranty_days=7)
        now = datetime.now(timezone.utc).replace(microsecond=0)

        unpaid_repo = StoreRepository(self.db_path)
        unpaid_repo.create_pending_account_order_and_reserve(
            order_id="ORD-TIMEOUT-UNPAID",
            telegram_user_id=1,
            username="Test User",
            product_code="GEMINI",
            product_name="GEMINI AI",
            package_name="Gemini AI Pro",
            quantity=1,
            unit_price_vnd=70000,
            total_vnd=70000,
            created_at=now.isoformat(),
            expire_at=(now + timedelta(minutes=5)).isoformat(),
        )
        with closing(sqlite3.connect(self.db_path)) as connection, connection:
            connection.execute(
                "UPDATE inventory_items SET reserved_at = ? WHERE reserved_order_id = ?",
                ((now - timedelta(minutes=6)).isoformat(), "ORD-TIMEOUT-UNPAID"),
            )
        released = unpaid_repo.release_expired_reservations(5)
        self.assertEqual(released, 1)
        state = snapshot_sales_state(self.db_path, "GEMINI", order_id="ORD-TIMEOUT-UNPAID")
        self.assertEqual(state.available_count, 8)
        self.assertEqual(state.reserved_count, 0)
        self.assertIn(state.order_status, {"expired", "cancelled", "pending", "reserved", "paid", ""})

        self._clear_inventory_for_codes("GEMINI")
        self._import_workbook("GEMINI", [f"gem{i}@example.com|pass" for i in range(1, 9)], price_vnd=70000, warranty_days=7)
        paid_repo = StoreRepository(self.db_path)
        paid_repo.create_pending_account_order_and_reserve(
            order_id="ORD-TIMEOUT-PAID",
            telegram_user_id=1,
            username="Test User",
            product_code="GEMINI",
            product_name="GEMINI AI",
            package_name="Gemini AI Pro",
            quantity=1,
            unit_price_vnd=70000,
            total_vnd=70000,
            created_at=now.isoformat(),
            expire_at=(now + timedelta(minutes=5)).isoformat(),
        )
        paid_repo.mark_account_order_paid_for_fulfillment("ORD-TIMEOUT-PAID")
        with closing(sqlite3.connect(self.db_path)) as connection, connection:
            connection.execute(
                "UPDATE inventory_items SET reserved_at = ? WHERE reserved_order_id = ?",
                ((now - timedelta(minutes=6)).isoformat(), "ORD-TIMEOUT-PAID"),
            )
        self.assertEqual(paid_repo.release_expired_reservations(5), 0)
        state = snapshot_sales_state(self.db_path, "GEMINI", order_id="ORD-TIMEOUT-PAID")
        self.assertEqual(state.available_count, 7)
        self.assertEqual(state.reserved_count, 1)
        delivered = paid_repo.deliver_reserved_items("ORD-TIMEOUT-PAID")
        self.assertEqual(len(delivered), 1)
        state = snapshot_sales_state(self.db_path, "GEMINI", order_id="ORD-TIMEOUT-PAID")
        self.assertEqual(state.available_count, 7)
        self.assertEqual(state.reserved_count, 0)
        self.assertEqual(state.delivered_count, 1)

    def test_replace_import_with_reserved_item_preserves_reserved_inventory(self) -> None:
        self._clear_inventory_for_codes("GEMINI")
        self._import_workbook("GEMINI", [f"gem{i}@example.com|pass" for i in range(1, 9)], price_vnd=70000, warranty_days=7)
        self._reserve_gemini_quantity(1, order_id="ORD-REPLACE-1")
        state_before = snapshot_sales_state(self.db_path, "GEMINI", order_id="ORD-REPLACE-1")
        self.assertEqual(state_before.available_count, 7)
        self.assertEqual(state_before.reserved_count, 1)

        replacement = Path(self.temp_dir.name) / "replacement.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "GEMINI"
        sheet.append(REQUIRED_COLUMNS)
        for index in range(5):
            sheet.append(["GEMINI", "account", "Gemini AI Pro", "account", "30D", 70000, 7, f"replace{index}@example.com|pass", "", 1])
        workbook.save(replacement)
        workbook.close()
        report = import_inventory(replacement, self.db_path, mode="replace")
        self.assertEqual(report["stock"], {"GEMINI": 5})

        state_after = snapshot_sales_state(self.db_path, "GEMINI", order_id="ORD-REPLACE-1")
        self.assertEqual(state_after.available_count, 5)
        self.assertEqual(state_after.reserved_count, 1)
        self.assertEqual(state_after.disabled_count, 7)
        validate_stock_invariant(state_after)

    def test_log_sales_state_reports_expected_fields(self) -> None:
        self._clear_inventory_for_codes("GEMINI")
        self._import_workbook("GEMINI", [f"gem{i}@example.com|pass" for i in range(1, 3)], price_vnd=70000, warranty_days=7)
        payload = log_sales_state(
            "diagnostic",
            "GEMINI",
            callback_data="qty:GEMINI AI:GEMINI:1",
            package_id="GEMINI",
            quantity=1,
            database_path=self.db_path,
        )
        self.assertEqual(payload["product_code"], "GEMINI")
        self.assertEqual(payload["available_count"], 2)
        self.assertTrue(payload["can_reserve"])
        self.assertEqual(payload["expected_available_after"], 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
